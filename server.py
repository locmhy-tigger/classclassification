# -*- coding: utf-8 -*-
import os
import glob
import re
import sys
import threading
import webbrowser
import queue
import tkinter as tk
from tkinter import filedialog
from flask import Flask, jsonify, request, send_from_directory
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

app = Flask(__name__, static_folder='dist', static_url_path='')

# Tkinter initialization
root = tk.Tk()
root.withdraw()
root.attributes('-topmost', True)

# Thread-safe dialog queue
dialog_queue = queue.Queue()

selected_paths = {
    'source_dir': r"G:\\我的雲端硬碟\\考評\\2025-2026",
    'output_path': r"G:\\我的雲端硬碟\\考評\\2025-2026\\25-26_全校分班分組_Generated.xlsx"
}

def find_file_by_pattern(directory, pattern):
    search_pattern = os.path.join(directory, pattern)
    matches = glob.glob(search_pattern)
    if matches:
        return matches[0]
    search_pattern_sub = os.path.join(directory, "*", pattern)
    matches_sub = glob.glob(search_pattern_sub)
    if matches_sub:
        return matches_sub[0]
    return None

def get_source_files(source_dir):
    files = {
        'student_list': find_file_by_pattern(source_dir, "*全校分班分組*(4Jul2025)*.xlsx") or find_file_by_pattern(source_dir, "*全校分班分組*.xlsx") or find_file_by_pattern(source_dir, "*student*info*.xlsx"),
        'annual_results': find_file_by_pattern(source_dir, "*年終*Result*.xlsx") or find_file_by_pattern(source_dir, "*年終*Result*.xls") or find_file_by_pattern(source_dir, "*年終成績*.xlsx"),
        'promotion_list': find_file_by_pattern(source_dir, "*升留人數及名單*.xlsx") or find_file_by_pattern(source_dir, "*升留*.xlsx") or find_file_by_pattern(source_dir, "*升留名單*.xlsx"),
        'different_class': find_file_by_pattern(source_dir, "*不同班建議*.xlsx") or find_file_by_pattern(source_dir, "*不同班*.xlsx"),
        'at_marks': find_file_by_pattern(source_dir, "*AT_Marks*.xlsx") or find_file_by_pattern(source_dir, "*AT*.xlsx") or find_file_by_pattern(source_dir, "*學前成績*.xlsx"),
        'electives': find_file_by_pattern(source_dir, "*選科及*重讀*.xlsx") or find_file_by_pattern(source_dir, "*選科*.xlsx")
    }
    return files

@app.route('/')
def serve_index():
    return send_from_directory(app.static_folder, 'index.html')

@app.route('/api/paths', methods=['GET'])
def get_paths():
    return jsonify(selected_paths)

@app.route('/api/select_folder', methods=['POST'])
def select_folder():
    res_q = queue.Queue()
    dialog_queue.put(('select_folder', res_q))
    folder = res_q.get()
    
    if folder:
        folder = os.path.abspath(folder)
        selected_paths['source_dir'] = folder
        selected_paths['output_path'] = os.path.join(folder, "Generated_全校分班分組.xlsx")
        return jsonify({'status': 'success', 'source_dir': folder, 'output_path': selected_paths['output_path']})
    return jsonify({'status': 'cancelled'})

@app.route('/api/select_output', methods=['POST'])
def select_output():
    res_q = queue.Queue()
    dialog_queue.put(('select_output', res_q))
    file_path = res_q.get()
    
    if file_path:
        file_path = os.path.abspath(file_path)
        selected_paths['output_path'] = file_path
        return jsonify({'status': 'success', 'output_path': file_path})
    return jsonify({'status': 'cancelled'})

@app.route('/api/load_source_status', methods=['GET'])
def load_source_status():
    source_dir = selected_paths['source_dir']
    if not source_dir or not os.path.exists(source_dir):
        return jsonify({'status': 'error', 'message': '請先選擇來源目錄'})
    
    files = get_source_files(source_dir)
    status = {}
    for key, path in files.items():
        status[key] = {
            'resolved': path is not None,
            'filename': os.path.basename(path) if path else '未找到符合的檔案',
            'path': path or ''
        }
    return jsonify(status)

@app.route('/api/load_initial_data', methods=['GET'])
def load_initial_data():
    source_dir = selected_paths['source_dir']
    if not source_dir or not os.path.exists(source_dir):
        return jsonify({'status': 'error', 'message': '請先選擇來源目錄'})
    
    files = get_source_files(source_dir)
    
    # 1. Load relation constraints (SEN / Different Class Suggestions)
    relations = []
    if files['different_class'] and os.path.exists(files['different_class']):
        try:
            df_diff = pd.read_excel(files['different_class'])
            for idx, row in df_diff.iterrows():
                cls1 = row.iloc[0]
                name1 = row.iloc[2]
                name2 = row.iloc[6]
                if pd.notna(name1) and pd.notna(name2) and str(name1).strip() != '' and str(name2).strip() != '' and str(name1).strip() != '姓名':
                    relations.append({
                        'student1': str(name1).strip().replace('_x000F_', '').replace('_x000E_', ''),
                        'student2': str(name2).strip().replace('_x000F_', '').replace('_x000E_', ''),
                        'reason': f"SEN / 訓輔不同班建議 (來源: {str(cls1).strip() if pd.notna(cls1) else ''})"
                    })
            
            # Search for extra text rules (e.g. 2B 蔡智康 and 朱霖)
            for idx, row in df_diff.iterrows():
                row_str = " ".join([str(x) for x in row.tolist() if pd.notna(x)])
                if 'and' in row_str and ('different' in row_str.lower() or 'suggestion' in row_str.lower() or 'TSOI' in row_str):
                    if '蔡智康' in row_str and '朱霖' in row_str:
                        relations.append({
                            'student1': '蔡智康',
                            'student2': '朱霖',
                            'reason': 'SEN 建議不同班 (蔡智康 & 朱霖)'
                        })
        except Exception as e:
            print(f"Error parsing different class recommendations: {e}")

    # 2. Load repeaters list
    repeaters = []
    if files['promotion_list'] and os.path.exists(files['promotion_list']):
        try:
            df_rep = pd.read_excel(files['promotion_list'], sheet_name='留級', skiprows=2, header=None)
            rep_groups = [
                ('S1', 0, 1, 2),
                ('S2', 4, 5, 6),
                ('S3', 8, 9, 10),
                ('S4', 12, 13, 14),
                ('S5', 16, 17, 18)
            ]
            for grade, c_cls, c_no, c_name in rep_groups:
                for idx, row in df_rep.iterrows():
                    name = row[c_name]
                    prev_cls = row[c_cls]
                    if pd.notna(name) and str(name).strip() != '' and str(name).strip() != '姓名':
                        clean_name = str(name).strip().replace('_x000F_', '').replace('_x000E_', '')
                        repeaters.append({
                            'name': clean_name,
                            'prev_class': str(prev_cls).strip(),
                            'grade': grade,
                            'status': 'Repeater',
                            'target_class': 'Auto'
                        })
        except Exception as e:
            print(f"Error parsing repeaters list: {e}")
            
    return jsonify({
        'relations': relations,
        'repeaters': repeaters,
        'files': {k: os.path.basename(v) if v else None for k, v in files.items()}
    })

@app.route('/api/process', methods=['POST'])
def process_data():
    data = request.json or {}
    source_dir = selected_paths['source_dir']
    output_path = selected_paths['output_path']
    config = data.get('config', {})
    manual_placements = data.get('manual_placements', {})
    
    if not source_dir or not os.path.exists(source_dir):
        return jsonify({'status': 'error', 'message': '無效的來源目錄'})
        
    files = get_source_files(source_dir)
    
    try:
        # Load last year's student list for mapping static properties (English Name, Gender, House, Logins)
        df_old_list = None
        if files['student_list'] and os.path.exists(files['student_list']):
            df_old_list = pd.read_excel(files['student_list'], sheet_name='All')
            df_old_list['Name_clean'] = df_old_list['中文姓名'].astype(str).str.strip().str.replace('_x000F_', '').str.replace('_x000E_', '')
            df_old_list['Reg_No_clean'] = df_old_list['學生註冊編號'].astype(str).str.strip()

        # Build name-gender and name-house map from old list & reference files if they exist to be 100% accurate
        name_gender_map = {}
        name_house_map = {}
        if df_old_list is not None:
            for idx, row in df_old_list.iterrows():
                name = row['Name_clean']
                gender = row['性別']
                house = row['社別']
                if pd.notna(name):
                    if pd.notna(gender): name_gender_map[name] = str(gender).strip()
                    if pd.notna(house): name_house_map[name] = str(house).strip()

        ref_path = os.path.join(source_dir, "25-26_全校分班分組_(24Jul2025) Final.xlsx")
        if os.path.exists(ref_path):
            try:
                ref_xls = pd.ExcelFile(ref_path)
                for s in ref_xls.sheet_names:
                    if s.startswith('S'):
                        ref_df = ref_xls.parse(s)
                        if '中文姓名' in ref_df.columns and '性別' in ref_df.columns:
                            for idx, row in ref_df.iterrows():
                                n = str(row['中文姓名']).strip().replace('_x000F_', '').replace('_x000E_', '')
                                g = row['性別']
                                h = row['社別'] if '社別' in ref_df.columns else None
                                if pd.notna(n) and n != 'nan' and n != '姓名':
                                    if pd.notna(g): name_gender_map[n] = str(g).strip()
                                    if pd.notna(h): name_house_map[n] = str(h).strip()
            except Exception as e:
                print(f"Error scanning reference file for genders: {e}")

        # Helper to pull old info
        def pull_student_info(reg_no=None, name=None):
            info = {
                '學生註冊編號': '',
                '英文姓名': '', '性別': '', '社別': '', '學習支援': '', '備忘錄': '',
                '選修1': '', '選修2': '', '選修3 / 應用學習': '',
                'eClass imail:': '', 'G-Suite login:': '', 'O365 Login:': '', 'Teams搜索': '', '教城帳戶': ''
            }
            if df_old_list is not None:
                match = pd.DataFrame()
                if reg_no:
                    match = df_old_list[df_old_list['Reg_No_clean'] == str(reg_no).strip()]
                if len(match) == 0 and name:
                    match = df_old_list[df_old_list['Name_clean'] == str(name).strip()]
                
                if len(match) > 0:
                    for col in info.keys():
                        if col in match.columns:
                            val = match.iloc[0][col]
                            info[col] = val if pd.notna(val) else ''
            
            if name and name in name_gender_map:
                info['性別'] = name_gender_map[name]
            if name and name in name_house_map:
                info['社別'] = name_house_map[name]
            return info

        # Load annual results
        xls_marks = pd.ExcelFile(files['annual_results'])
        marks_sheets = {}
        for s in xls_marks.sheet_names:
            df_m = xls_marks.parse(s)
            if '學生姓名' not in df_m.columns:
                for idx, row in df_m.iterrows():
                    if '學生姓名' in row.values:
                        df_m.columns = row.values
                        df_m = df_m.iloc[idx+1:]
                        break
            if '學生姓名' in df_m.columns:
                df_m['Name_clean'] = df_m['學生姓名'].astype(str).str.strip().str.replace('_x000F_', '').str.replace('_x000E_', '')
                marks_sheets[s] = df_m

        placed_students = {f'S{i}': [] for i in range(1, 7)}

        # Load repeaters list
        promotion_repeaters = []
        if files['promotion_list'] and os.path.exists(files['promotion_list']):
            df_rep = pd.read_excel(files['promotion_list'], sheet_name='留級', skiprows=2, header=None)
            rep_groups = [
                ('S1', 0, 1, 2),
                ('S2', 4, 5, 6),
                ('S3', 8, 9, 10),
                ('S4', 12, 13, 14),
                ('S5', 16, 17, 18)
            ]
            for grade, c_cls, c_no, c_name in rep_groups:
                for idx, row in df_rep.iterrows():
                    name = row.iloc[c_name]
                    prev_cls = row.iloc[c_cls]
                    if pd.notna(name) and str(name).strip() != '' and str(name).strip() != '姓名':
                        clean_name = str(name).strip().replace('_x000F_', '').replace('_x000E_', '')
                        promotion_repeaters.append({
                            'name': clean_name,
                            'prev_class': str(prev_cls).strip(),
                            'grade': grade,
                            'status': 'Repeater'
                        })

        # =========================================================================
        # S1 Placement
        # =========================================================================
        s1_students = []
        if files['at_marks'] and os.path.exists(files['at_marks']):
            df_at_raw = pd.read_excel(files['at_marks'], sheet_name='總分')
            h_row = 1
            for idx, r in df_at_raw.iterrows():
                r_str = [str(x) for x in r.tolist()]
                if any('Reg. No.' in x or 'STRN' in x for x in r_str):
                    h_row = idx
                    break
            df_at = pd.read_excel(files['at_marks'], sheet_name='總分', header=h_row + 1)
            df_at = df_at[df_at['Reg. No.'].notna()]
            df_at['Name_clean'] = df_at['中文姓名'].astype(str).str.strip().str.replace('_x000F_', '').str.replace('_x000E_', '')
            df_at['總分_num'] = pd.to_numeric(df_at['總分'], errors='coerce')
            df_at['Reg_No_clean'] = df_at['Reg. No.'].astype(str).str.strip()
            df_at = df_at.sort_values(by=['總分_num', 'Reg_No_clean'], ascending=[False, True])
            
            s1_top_count = int(config.get('s1_top_count', 35))
            s1_top_class = config.get('s1_top_class', '1A')
            s1_other_pattern = config.get('s1_other_pattern', 'gender_snake')
            
            s1_repeaters_from_list = [r for r in promotion_repeaters if r['grade'] == 'S1']
            
            top_df = df_at.iloc[:s1_top_count].copy()
            other_df = df_at.iloc[s1_top_count:].copy()
            
            for idx, row in top_df.iterrows():
                name = row['Name_clean']
                info = pull_student_info(reg_no=row['Reg_No_clean'], name=name)
                s1_students.append({
                    'Name': name,
                    'Reg_No': row['Reg_No_clean'],
                    'Class': s1_top_class,
                    'Avg': row['總分_num'],
                    'Rank': row['Rank'],
                    'Chinese': row['中文'], 'English': row['英文'], 'Math': row['數學'], 'LS': '',
                    'EnglishName': row['英文姓名'] if pd.notna(row['英文姓名']) else info['英文姓名'],
                    'Gender': info['性別'] or 'M',
                    'House': info['社別'] or '',
                    'Support': info['學習支援'] or '',
                    'Remark': info['備忘錄'] or '',
                    'IsRepeater': False
                })
                
            other_students_list = []
            for idx, row in other_df.iterrows():
                name = row['Name_clean']
                info = pull_student_info(reg_no=row['Reg_No_clean'], name=name)
                other_students_list.append({
                    'Name': name,
                    'Reg_No': row['Reg_No_clean'],
                    'Avg': row['總分_num'],
                    'Rank': row['Rank'],
                    'Chinese': row['中文'], 'English': row['英文'], 'Math': row['數學'], 'LS': '',
                    'EnglishName': row['英文姓名'] if pd.notna(row['英文姓名']) else info['英文姓名'],
                    'Gender': info['性別'] or 'M',
                    'House': info['社別'] or '',
                    'Support': info['學習支援'] or '',
                    'Remark': info['備忘錄'] or '',
                    'IsRepeater': False
                })
                
            for rep in s1_repeaters_from_list:
                name = rep['name']
                info = pull_student_info(name=name)
                other_students_list.append({
                    'Name': name,
                    'Reg_No': info['學生註冊編號'] or '',
                    'Avg': '', 'Rank': '', 'Chinese': '', 'English': '', 'Math': '', 'LS': '',
                    'EnglishName': info['英文姓名'],
                    'Gender': info['性別'] or 'F',
                    'House': info['社別'] or '',
                    'Support': info['學習支援'] or '',
                    'Remark': 'R',
                    'IsRepeater': True
                })

            auto_students = []
            for s in other_students_list:
                if s['Name'] in manual_placements:
                    s['Class'] = manual_placements[s['Name']]
                    s1_students.append(s)
                else:
                    auto_students.append(s)
                    
            classes = ['1B', '1C', '1D']
            if s1_other_pattern == 'gender_snake':
                females = [s for s in auto_students if s['Gender'] == 'F']
                males = [s for s in auto_students if s['Gender'] == 'M']
                females = sorted(females, key=lambda x: x['Avg'] if isinstance(x['Avg'], (int, float)) else -1, reverse=True)
                males = sorted(males, key=lambda x: x['Avg'] if isinstance(x['Avg'], (int, float)) else -1, reverse=True)
                
                for i, s in enumerate(females):
                    cycle = i % 6
                    if cycle == 0: s['Class'] = '1B'
                    elif cycle == 1: s['Class'] = '1C'
                    elif cycle == 2: s['Class'] = '1D'
                    elif cycle == 3: s['Class'] = '1D'
                    elif cycle == 4: s['Class'] = '1C'
                    elif cycle == 5: s['Class'] = '1B'
                    s1_students.append(s)
                for i, s in enumerate(males):
                    cycle = i % 6
                    if cycle == 0: s['Class'] = '1B'
                    elif cycle == 1: s['Class'] = '1C'
                    elif cycle == 2: s['Class'] = '1D'
                    elif cycle == 3: s['Class'] = '1D'
                    elif cycle == 4: s['Class'] = '1C'
                    elif cycle == 5: s['Class'] = '1B'
                    s1_students.append(s)
            else:
                auto_students = sorted(auto_students, key=lambda x: x['Avg'] if isinstance(x['Avg'], (int, float)) else -1, reverse=True)
                for i, s in enumerate(auto_students):
                    s['Class'] = classes[i % len(classes)]
                    s1_students.append(s)

            new_students = [s for s in s1_students if s['Reg_No'].startswith('C5') and s['House'] == '']
            new_students = sorted(new_students, key=lambda x: x['Reg_No'])
            houses = ['紅社', '黃社', '藍社', '綠社']
            for i, s in enumerate(new_students):
                s['House'] = houses[i % len(houses)]

            placed_students['S1'] = s1_students

        # =========================================================================
        # S2 Placement
        # =========================================================================
        s2_students = []
        df_s1_marks = marks_sheets.get('S1')
        if df_s1_marks is not None:
            df_s1_marks['Rank_num'] = pd.to_numeric(df_s1_marks['級別名次'], errors='coerce')
            s1_promoted = df_s1_marks.sort_values('Rank_num').copy()
            s1_rep_names = [r['name'] for r in promotion_repeaters if r['grade'] == 'S1']
            s1_promoted = s1_promoted[~s1_promoted['Name_clean'].isin(s1_rep_names)]
            
            s2_top_count = int(config.get('s2_top_count', 34))
            s2_top_class = config.get('s2_top_class', '2A')
            
            top_df = s1_promoted.iloc[:s2_top_count].copy()
            other_df = s1_promoted.iloc[s2_top_count:].copy()
            
            for idx, row in top_df.iterrows():
                name = row['Name_clean']
                info = pull_student_info(name=name)
                s2_students.append({
                    'Name': name, 'Reg_No': info['學生註冊編號'] or '', 'Class': s2_top_class,
                    'Avg': row['平均分'], 'Rank': row['級別名次'],
                    'Chinese': row['中文'], 'English': row['英文'], 'Math': row['數學'], 'LS': row['公經社'] if '公經社' in row else '',
                    'EnglishName': info['英文姓名'], 'Gender': info['性別'], 'House': info['社別'],
                    'Support': info['學習支援'], 'Remark': info['備忘錄'], 'IsRepeater': False
                })
                
            classes = ['2B', '2C', '2D']
            other_students_list = []
            for idx, row in other_df.iterrows():
                name = row['Name_clean']
                info = pull_student_info(name=name)
                other_students_list.append({
                    'Name': name, 'Reg_No': info['學生註冊編號'] or '',
                    'Avg': row['平均分'], 'Rank': row['級別名次'],
                    'Chinese': row['中文'], 'English': row['英文'], 'Math': row['數學'], 'LS': row['公經社'] if '公經社' in row else '',
                    'EnglishName': info['英文姓名'], 'Gender': info['性別'], 'House': info['社別'],
                    'Support': info['學習支援'], 'Remark': info['備忘錄'], 'IsRepeater': False
                })
                
            s2_repeaters = [r for r in promotion_repeaters if r['grade'] == 'S2']
            for rep in s2_repeaters:
                name = rep['name']
                info = pull_student_info(name=name)
                other_students_list.append({
                    'Name': name, 'Reg_No': info['學生註冊編號'] or '',
                    'Avg': '', 'Rank': '', 'Chinese': '', 'English': '', 'Math': '', 'LS': '',
                    'EnglishName': info['英文姓名'], 'Gender': info['性別'], 'House': info['社別'],
                    'Support': info['學習支援'], 'Remark': 'R', 'IsRepeater': True
                })
                
            auto_students = []
            for s in other_students_list:
                if s['Name'] in manual_placements:
                    s['Class'] = manual_placements[s['Name']]
                    s2_students.append(s)
                elif s['IsRepeater'] and config.get('s2_repeater_placement') == 'previous':
                    prev_let = str(s['Class']).strip()[-1] if 'Class' in s else 'B'
                    s['Class'] = '2' + prev_let
                    s2_students.append(s)
                else:
                    auto_students.append(s)
                    
            auto_students = sorted(auto_students, key=lambda x: x['Rank'] if isinstance(x['Rank'], (int, float)) else 999)
            for i, s in enumerate(auto_students):
                s['Class'] = classes[i % len(classes)]
                s2_students.append(s)
                
            placed_students['S2'] = s2_students

        # =========================================================================
        # S3 Placement
        # =========================================================================
        s3_students = []
        df_s2_marks = marks_sheets.get('S2')
        if df_s2_marks is not None:
            df_s2_marks['Rank_num'] = pd.to_numeric(df_s2_marks['級別名次'], errors='coerce')
            s2_promoted = df_s2_marks.sort_values('Rank_num').copy()
            s2_rep_names = [r['name'] for r in promotion_repeaters if r['grade'] == 'S2']
            s2_promoted = s2_promoted[~s2_promoted['Name_clean'].isin(s2_rep_names)]
            
            s3_top_count = int(config.get('s3_top_count', 35))
            s3_top_class = config.get('s3_top_class', '3A')
            
            top_df = s2_promoted.iloc[:s3_top_count].copy()
            other_df = s2_promoted.iloc[s3_top_count:].copy()
            
            for idx, row in top_df.iterrows():
                name = row['Name_clean']
                info = pull_student_info(name=name)
                s3_students.append({
                    'Name': name, 'Reg_No': info['學生註冊編號'] or '', 'Class': s3_top_class,
                    'Avg': row['平均分'], 'Rank': row['級別名次'],
                    'Chinese': row['中文'], 'English': row['英文'], 'Math': row['數學'], 'LS': row['生社'] if '生社' in row else '',
                    'EnglishName': info['英文姓名'], 'Gender': info['性別'], 'House': info['社別'],
                    'Support': info['學習支援'], 'Remark': info['備忘錄'], 'IsRepeater': False
                })
                
            other_students_list = []
            for idx, row in other_df.iterrows():
                name = row['Name_clean']
                info = pull_student_info(name=name)
                prev_class = str(row['班別']).strip()
                target_class = ''
                if prev_class != '2A':
                    target_class = '3' + prev_class[-1]
                    
                other_students_list.append({
                    'Name': name, 'Reg_No': info['學生註冊編號'] or '',
                    'Avg': row['平均分'], 'Rank': row['級別名次'],
                    'Chinese': row['中文'], 'English': row['英文'], 'Math': row['數學'], 'LS': row['生社'] if '生社' in row else '',
                    'EnglishName': info['英文姓名'], 'Gender': info['性別'], 'House': info['社別'],
                    'Support': info['學習支援'], 'Remark': info['備忘錄'], 'IsRepeater': False,
                    'Prev_Class': prev_class,
                    'Class': target_class
                })
                
            s3_repeaters = [r for r in promotion_repeaters if r['grade'] == 'S3']
            for rep in s3_repeaters:
                name = rep['name']
                info = pull_student_info(name=name)
                other_students_list.append({
                    'Name': name, 'Reg_No': info['學生註冊編號'] or '',
                    'Avg': '', 'Rank': '', 'Chinese': '', 'English': '', 'Math': '', 'LS': '',
                    'EnglishName': info['英文姓名'], 'Gender': info['性別'], 'House': info['社別'],
                    'Support': info['學習支援'], 'Remark': 'R', 'IsRepeater': True,
                    'Prev_Class': rep['prev_class'],
                    'Class': ''
                })
                
            auto_remaining = []
            for s in other_students_list:
                if s['Name'] in manual_placements:
                    s['Class'] = manual_placements[s['Name']]
                    s3_students.append(s)
                elif s['Class'] != '':
                    s3_students.append(s)
                else:
                    auto_remaining.append(s)
                    
            classes = ['3B', '3C', '3D']
            for i, s in enumerate(auto_remaining):
                s['Class'] = classes[i % len(classes)]
                s3_students.append(s)
                
            placed_students['S3'] = s3_students

        # =========================================================================
        # S4 Placement
        # =========================================================================
        s4_students = []
        df_s3_marks = marks_sheets.get('S3')
        
        electives_dict = {}
        if files['electives'] and os.path.exists(files['electives']):
            df_el_raw = pd.read_excel(files['electives'], sheet_name='25-26 中四', header=None)
            blocks = [(0, '3A'), (7, '3B'), (15, '3C'), (23, '3D'), (31, 'Repeaters')]
            for col_start, b_name in blocks:
                for idx in range(1, len(df_el_raw)):
                    name = df_el_raw.iloc[idx, col_start + 2]
                    opt1 = df_el_raw.iloc[idx, col_start + 3]
                    opt2 = df_el_raw.iloc[idx, col_start + 4]
                    opt3 = df_el_raw.iloc[idx, col_start + 5]
                    if pd.notna(name) and str(name).strip() != '' and str(name).strip() != '姓名':
                        c_name = str(name).strip().replace('_x000F_', '').replace('_x000E_', '')
                        electives_dict[c_name] = {
                            'X1': str(opt1).strip() if pd.notna(opt1) else '',
                            'X2': str(opt2).strip() if pd.notna(opt2) else '',
                            'X3': str(opt3).strip() if pd.notna(opt3) else ''
                        }

        if df_s3_marks is not None:
            df_s3_marks['Rank_num'] = pd.to_numeric(df_s3_marks['級別名次'], errors='coerce')
            s3_promoted = df_s3_marks.sort_values('Rank_num').copy()
            s3_rep_names = [r['name'] for r in promotion_repeaters if r['grade'] == 'S3']
            s3_promoted = s3_promoted[~s3_promoted['Name_clean'].isin(s3_rep_names)]
            
            students_s4_pool = []
            for idx, row in s3_promoted.iterrows():
                name = row['Name_clean']
                info = pull_student_info(name=name)
                el = electives_dict.get(name, {'X1': '', 'X2': '', 'X3': ''})
                
                chi = pd.to_numeric(row['中文'], errors='coerce') or 0
                ls = pd.to_numeric(row['生社'], errors='coerce') or 0
                chi_ls = chi + ls
                
                is_x3 = el['X3'] in ['M1', '生物']
                
                students_s4_pool.append({
                    'Name': name, 'Reg_No': info['學生註冊編號'] or '',
                    'Avg': row['平均分'], 'Rank': row['級別名次'], 'Rank_num': row['Rank_num'],
                    'Chinese': row['中文'], 'English': row['英文'], 'Math': row['數學'], 'LS': row['生社'],
                    'EnglishName': info['英文姓名'], 'Gender': info['性別'], 'House': info['社別'],
                    'Support': info['學習支援'], 'Remark': info['備忘錄'], 'IsRepeater': False,
                    'X1': el['X1'], 'X2': el['X2'], 'X3': el['X3'],
                    'is_X3': is_x3, 'CHI_LS': chi_ls
                })
                
            s4_x3_top_count = int(config.get('s4_x3_top_count', 33))
            s4_x3_class = config.get('s4_x3_class', '4D')
            
            x3_students = [s for s in students_s4_pool if s['is_X3']]
            x3_students = sorted(x3_students, key=lambda x: x['Rank_num'])
            
            top_x3 = x3_students[:s4_x3_top_count]
            remaining_x3 = x3_students[s4_x3_top_count:]
            
            for s in top_x3:
                s['Class'] = s4_x3_class
                s4_students.append(s)
                
            for s in remaining_x3:
                s['Class'] = '4C'
                s4_students.append(s)
                
            target_4c_size = int(config.get('s4_4c_size', 31))
            filled_4c_count = len(remaining_x3)
            spots_to_fill = max(0, target_4c_size - filled_4c_count)
            
            placed_names = set([s['Name'] for s in s4_students])
            non_x3_students = [s for s in students_s4_pool if not s['Name'] in placed_names]
            non_x3_students = sorted(non_x3_students, key=lambda x: x['CHI_LS'], reverse=True)
            
            top_non_x3_for_4c = non_x3_students[:spots_to_fill]
            for s in top_non_x3_for_4c:
                s['Class'] = '4C'
                s4_students.append(s)
                
            placed_names = set([s['Name'] for s in s4_students])
            remaining_for_ab = [s for s in students_s4_pool if not s['Name'] in placed_names]
            remaining_for_ab = sorted(remaining_for_ab, key=lambda x: x['Rank_num'])
            
            s4_repeaters = [r for r in promotion_repeaters if r['grade'] == 'S4']
            s4_repeaters_list = []
            for rep in s4_repeaters:
                name = rep['name']
                info = pull_student_info(name=name)
                el = electives_dict.get(name, {'X1': '', 'X2': '', 'X3': ''})
                s4_repeaters_list.append({
                    'Name': name, 'Reg_No': info['學生註冊編號'] or '',
                    'Avg': '', 'Rank': '', 'Rank_num': 999, 'Chinese': '', 'English': '', 'Math': '', 'LS': '',
                    'EnglishName': info['英文姓名'], 'Gender': info['性別'], 'House': info['社別'],
                    'Support': info['學習支援'], 'Remark': 'R', 'IsRepeater': True,
                    'X1': el['X1'], 'X2': el['X2'], 'X3': el['X3'],
                    'is_X3': False, 'CHI_LS': 0, 'Prev_Class': rep['prev_class']
                })
                
            auto_remaining_ab = []
            for s in remaining_for_ab + s4_repeaters_list:
                if s['Name'] in manual_placements:
                    s['Class'] = manual_placements[s['Name']]
                    s4_students.append(s)
                else:
                    auto_remaining_ab.append(s)
                    
            ab_classes = ['4A', '4B']
            auto_remaining_ab = sorted(auto_remaining_ab, key=lambda x: x['Rank_num'])
            for i, s in enumerate(auto_remaining_ab):
                s['Class'] = ab_classes[i % len(ab_classes)]
                s4_students.append(s)
                
            placed_students['S4'] = s4_students

        # =========================================================================
        # S5 Placement
        # =========================================================================
        s5_students = []
        df_s4_marks = marks_sheets.get('S4')
        
        s5_rep_electives = {}
        if files['electives'] and os.path.exists(files['electives']):
            try:
                df_rep_el = pd.read_excel(files['electives'], sheet_name='25-26 中五重讀', header=None)
                for idx in range(2, len(df_rep_el)):
                    name = df_rep_el.iloc[idx, 3]
                    opt1 = df_rep_el.iloc[idx, 4]
                    opt2 = df_rep_el.iloc[idx, 5]
                    opt3 = df_rep_el.iloc[idx, 6]
                    apl = df_rep_el.iloc[idx, 7]
                    ret = df_rep_el.iloc[idx, 8]
                    if pd.notna(name) and str(name).strip() != '' and str(name).strip() != '姓名':
                        c_name = str(name).strip().replace('_x000F_', '').replace('_x000E_', '')
                        x3_val = str(opt3).strip() if pd.notna(opt3) else ''
                        if x3_val == '' and pd.notna(apl):
                            x3_val = str(apl).strip()
                        s5_rep_electives[c_name] = {
                            'X1': str(opt1).strip() if pd.notna(opt1) else '',
                            'X2': str(opt2).strip() if pd.notna(opt2) else '',
                            'X3': x3_val,
                            'Target_Class': str(df_rep_el.iloc[idx, 1]).strip() if pd.notna(df_rep_el.iloc[idx, 1]) else '5A'
                        }
            except Exception as e:
                print(f"Error parsing S5 repeater electives: {e}")

        if df_s4_marks is not None:
            s4_promoted = df_s4_marks.copy()
            s4_rep_names = [r['name'] for r in promotion_repeaters if r['grade'] == 'S4']
            s4_promoted = s4_promoted[~s4_promoted['Name_clean'].isin(s4_rep_names)]
            
            for idx, row in s4_promoted.iterrows():
                name = row['Name_clean']
                info = pull_student_info(name=name)
                prev_class = str(row.get('班別', '')).strip()
                target_class = '5' + prev_class[-1] if prev_class else '5A'

                if name in manual_placements:
                    target_class = manual_placements[name]

                s5_students.append({
                    'Name': name, 'Reg_No': info['學生註冊編號'] or '', 'Class': target_class,
                    'Avg': row.get('平均分', ''), 'Rank': row.get('級別名次', ''),
                    'Chinese': row.get('中文', ''), 'English': row.get('英文', ''), 'Math': row.get('數學', ''), 'LS': '',
                    'EnglishName': info['英文姓名'], 'Gender': info['性別'], 'House': info['社別'],
                    'Support': info['學習支援'], 'Remark': info['備忘錄'], 'IsRepeater': False,
                    'X1': info['選修1'], 'X2': info['選修2'], 'X3': info['選修3 / 應用學習']
                })

            s5_repeaters = [r for r in promotion_repeaters if r['grade'] == 'S5']
            for rep in s5_repeaters:
                name = rep['name']
                info = pull_student_info(name=name)
                rep_el = s5_rep_electives.get(name, {'X1': '', 'X2': '', 'X3': '', 'Target_Class': '5A'})
                
                target_class = rep_el['Target_Class']
                if name in manual_placements:
                    target_class = manual_placements[name]
                    
                s5_students.append({
                    'Name': name, 'Reg_No': info['學生註冊編號'] or '', 'Class': target_class,
                    'Avg': '', 'Rank': '', 'Chinese': '', 'English': '', 'Math': '', 'LS': '',
                    'EnglishName': info['英文姓名'], 'Gender': info['性別'], 'House': info['社別'],
                    'Support': info['學習支援'], 'Remark': 'R', 'IsRepeater': True,
                    'X1': rep_el['X1'] or info['選修1'], 'X2': rep_el['X2'] or info['選修2'], 'X3': rep_el['X3'] or info['選修3 / 應用學習']
                })
                
            placed_students['S5'] = s5_students

        # =========================================================================
        # S6 Placement
        # =========================================================================
        s6_students = []
        df_s5_marks = marks_sheets.get('S5')
        if df_s5_marks is not None:
            s5_promoted = df_s5_marks.copy()
            s5_rep_names = [r['name'] for r in promotion_repeaters if r['grade'] == 'S5']
            s5_promoted = s5_promoted[~s5_promoted['Name_clean'].isin(s5_rep_names)]
            
            for idx, row in s5_promoted.iterrows():
                name = row['Name_clean']
                info = pull_student_info(name=name)
                prev_class = str(row.get('班別', '')).strip()
                target_class = '6' + prev_class[-1] if prev_class else '6A'

                if name in manual_placements:
                    target_class = manual_placements[name]

                s6_students.append({
                    'Name': name, 'Reg_No': info['學生註冊編號'] or '', 'Class': target_class,
                    'Avg': row.get('平均分', ''), 'Rank': row.get('級別名次', ''),
                    'Chinese': row.get('中文', ''), 'English': row.get('英文', ''), 'Math': row.get('數學', ''), 'LS': '',
                    'EnglishName': info['英文姓名'], 'Gender': info['性別'], 'House': info['社別'],
                    'Support': info['學習支援'], 'Remark': info['備忘錄'], 'IsRepeater': False,
                    'X1': info['選修1'], 'X2': info['選修2'], 'X3': info['選修3 / 應用學習']
                })
                
            s6_repeaters = [r for r in promotion_repeaters if r['grade'] == 'S6']
            for rep in s6_repeaters:
                name = rep['name']
                info = pull_student_info(name=name)
                
                target_class = '6B'
                if name in manual_placements:
                    target_class = manual_placements[name]
                    
                s6_students.append({
                    'Name': name, 'Reg_No': info['學生註冊編號'] or '', 'Class': target_class,
                    'Avg': '', 'Rank': '', 'Chinese': '', 'English': '', 'Math': '', 'LS': '',
                    'EnglishName': info['英文姓名'], 'Gender': info['性別'], 'House': info['社別'],
                    'Support': info['學習支援'], 'Remark': 'R', 'IsRepeater': True,
                    'X1': info['選修1'], 'X2': info['選修2'], 'X3': info['選修3 / 應用學習']
                })
                
            placed_students['S6'] = s6_students

        # =========================================================================
        # Post-process sheets
        # =========================================================================
        final_placed = {}
        for grade, list_st in placed_students.items():
            df_g = pd.DataFrame(list_st)
            if len(df_g) > 0:
                df_g['EnglishName_sort'] = df_g['EnglishName'].astype(str).str.strip().str.upper()
                df_g['EnglishName_sort'] = df_g['EnglishName_sort'].map(lambda x: 'ZZZ' if x == '' or x == 'NAN' else x)
                df_g = df_g.sort_values(by=['Class', 'EnglishName_sort', 'Name'])
                df_g['25-26 班號'] = df_g.groupby('Class').cumcount() + 1
                
                for idx, row in df_g.iterrows():
                    name = row['Name']
                    reg = row['Reg_No']
                    info = pull_student_info(reg_no=reg, name=name)
                    for k, v in info.items():
                        if k not in df_g.columns or pd.isna(df_g.at[idx, k]) or str(df_g.at[idx, k]).strip() == '':
                            df_g.at[idx, k] = v
                
                df_g['25-26 班別'] = df_g['Class']
                df_g['平均分'] = df_g['Avg']
                df_g['級別名次'] = df_g['Rank']
                df_g['中文姓名'] = df_g['Name']
                df_g['學生註冊編號'] = df_g['Reg_No']
                df_g['性別'] = df_g['Gender']
                df_g['社別'] = df_g['House']
                df_g['學習支援'] = df_g['Support']
                df_g['備忘錄'] = df_g['Remark']
                
                if 'X1' in df_g.columns: df_g['選修1'] = df_g['X1']
                if 'X2' in df_g.columns: df_g['選修2'] = df_g['X2']
                if 'X3' in df_g.columns: df_g['選修3 / 應用學習'] = df_g['X3']
                
                final_placed[grade] = df_g
            else:
                final_placed[grade] = pd.DataFrame()

        # =========================================================================
        # Write Output Excel
        # =========================================================================
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        
        thin_border = Border(
            left=Side(style='thin', color='CCCCCC'),
            right=Side(style='thin', color='CCCCCC'),
            top=Side(style='thin', color='CCCCCC'),
            bottom=Side(style='thin', color='CCCCCC')
        )
        
        font_header = Font(name='新細明體', size=11, bold=True)
        font_body = Font(name='Times New Roman', size=11)
        font_chinese = Font(name='新細明體', size=11)
        font_summary = Font(name='Times New Roman', size=11, bold=True)
        
        align_center = Alignment(horizontal='center', vertical='center')
        align_left = Alignment(horizontal='left', vertical='center')
        
        headers_all = [
            '25-26 班別', '25-26 班號', '平均分', '級別名次', '中文', '英文', '數學', '公經社', 
            '中文姓名', '英文姓名', '學生註冊編號', '性別', '社別', '備忘錄', '學習支援', 
            '中文閱讀', '中文分組', '數學分組', '英文分組', '選修1', '選修2', '選修3 / 應用學習', 
            'OLE   分組', '宗教 分組', '中六中文提升班', '中六數學提升班', '中六英文提升班', 
            '中六 X1 退修班', '中六 X2 退修班', 'eClass imail:', 'G-Suite login:', 'O365 Login:', 
            'Teams搜索', '教城帳戶'
        ]
        
        class_positions = {}
        
        for grade in ['S6', 'S5', 'S4', 'S3', 'S2', 'S1']:
            ws = wb.create_sheet(title=grade)
            for col_idx, h in enumerate(headers_all, 1):
                cell = ws.cell(row=1, column=col_idx, value=h)
                cell.font = font_header
                cell.alignment = align_center
                cell.fill = PatternFill(start_color='EFEFEF', end_color='EFEFEF', fill_type='solid')
                cell.border = thin_border
                
            df_g = final_placed.get(grade, pd.DataFrame())
            row_idx = 2
            
            class_positions[grade] = {}
            current_class = None
            class_start_row = 2
            
            if len(df_g) > 0:
                for col in headers_all:
                    if col not in df_g.columns:
                        df_g[col] = ''
                        
                df_g = df_g[headers_all]
                
                for idx, row in df_g.iterrows():
                    cls = str(row['25-26 班別']).strip()
                    if cls != current_class:
                        if current_class is not None:
                            class_positions[grade][current_class] = (class_start_row, row_idx - 1)
                        current_class = cls
                        class_start_row = row_idx
                        
                    for col_idx, col_name in enumerate(headers_all, 1):
                        val = row[col_name]
                        if pd.isna(val) or val == 'nan':
                            val = ''
                        
                        cell = ws.cell(row=row_idx, column=col_idx, value=val)
                        cell.border = thin_border
                        
                        if col_name in ['中文姓名']:
                            cell.font = font_chinese
                            cell.alignment = align_left
                        elif col_name in ['英文姓名', '學生註冊編號', 'eClass imail:', 'G-Suite login:', 'O365 Login:', '教城帳戶']:
                            cell.font = font_body
                            cell.alignment = align_left
                        else:
                            cell.font = font_body
                            cell.alignment = align_center
                            
                    row_idx += 1
                
                if current_class is not None:
                    class_positions[grade][current_class] = (class_start_row, row_idx - 1)
                    
            summary_header_row = row_idx + 1
            ws.cell(row=summary_header_row, column=3, value='').border = thin_border
            
            ws.cell(row=summary_header_row + 1, column=4, value='人數').font = font_summary
            ws.cell(row=summary_header_row + 1, column=4, alignment=align_center)
            ws.cell(row=summary_header_row + 1, column=5, value='F').font = font_summary
            ws.cell(row=summary_header_row + 1, column=5, alignment=align_center)
            ws.cell(row=summary_header_row + 1, column=6, value='M').font = font_summary
            ws.cell(row=summary_header_row + 1, column=6, alignment=align_center)
            
            curr_sum_row = summary_header_row + 2
            
            grade_classes = sorted(list(class_positions[grade].keys()))
            if not grade_classes:
                grade_classes = [f'{grade[-1]}A', f'{grade[-1]}B', f'{grade[-1]}C', f'{grade[-1]}D']
                
            for cls in grade_classes:
                ws.cell(row=curr_sum_row, column=3, value=cls).font = font_summary
                ws.cell(row=curr_sum_row, column=3, alignment=align_center)
                
                total_formula = f'=COUNTIF($A$2:$A${row_idx-1}, "{cls}")'
                ws.cell(row=curr_sum_row, column=4, value=total_formula).font = font_summary
                ws.cell(row=curr_sum_row, column=4, alignment=align_center)
                
                start, end = class_positions[grade].get(cls, (2, 2))
                female_formula = f'=COUNTIF($L${start}:$L${end}, "F")'
                male_formula = f'=COUNTIF($L${start}:$L${end}, "M")'
                
                ws.cell(row=curr_sum_row, column=5, value=female_formula).font = font_summary
                ws.cell(row=curr_sum_row, column=5, alignment=align_center)
                
                ws.cell(row=curr_sum_row, column=6, value=male_formula).font = font_summary
                ws.cell(row=curr_sum_row, column=6, alignment=align_center)
                
                class_positions[grade][cls + '_sum_cells'] = {
                    'total': f'={grade}!D{curr_sum_row}',
                    'female': f'={grade}!E{curr_sum_row}',
                    'male': f'={grade}!F{curr_sum_row}'
                }
                
                curr_sum_row += 1
                
            for col in range(1, len(headers_all) + 1):
                col_letter = get_column_letter(col)
                ws.column_dimensions[col_letter].width = 15

        # =========================================================================
        # Write Sheet '人數'
        # =========================================================================
        ws_num = wb.create_sheet(title='人數')
        ws_num.cell(row=1, column=3, value='人數').font = font_header
        ws_num.cell(row=1, column=4, value='F').font = font_header
        ws_num.cell(row=1, column=5, value='M').font = font_header
        
        curr_row = 2
        for grade in ['S1', 'S2', 'S3', 'S4', 'S5', 'S6']:
            g_num = grade[-1]
            for letter in ['A', 'B', 'C', 'D']:
                cls = g_num + letter
                ws_num.cell(row=curr_row, column=2, value=cls).font = font_body
                ws_num.cell(row=curr_row, column=2, alignment=align_center)
                
                if grade in class_positions and cls + '_sum_cells' in class_positions[grade]:
                    cells = class_positions[grade][cls + '_sum_cells']
                    ws_num.cell(row=curr_row, column=3, value=cells['total']).font = font_body
                    ws_num.cell(row=curr_row, column=4, value=cells['female']).font = font_body
                    ws_num.cell(row=curr_row, column=5, value=cells['male']).font = font_body
                else:
                    ws_num.cell(row=curr_row, column=3, value=0).font = font_body
                    ws_num.cell(row=curr_row, column=4, value=0).font = font_body
                    ws_num.cell(row=curr_row, column=5, value=0).font = font_body
                
                curr_row += 1
                
        wb.save(output_path)
        
        response_stats = {}
        for grade in ['S1', 'S2', 'S3', 'S4', 'S5', 'S6']:
            df_g = final_placed.get(grade, pd.DataFrame())
            response_stats[grade] = {}
            if len(df_g) > 0:
                for cls in df_g['25-26 班別'].unique():
                    cls_df = df_g[df_g['25-26 班別'] == cls]
                    response_stats[grade][cls] = {
                        'total': len(cls_df),
                        'female': int((cls_df['性別'] == 'F').sum()),
                        'male': int((cls_df['性別'] == 'M').sum()),
                        'students': cls_df[['中文姓名', '英文姓名', '性別', '學生註冊編號', '備忘錄', '25-26 班號']].to_dict('records')
                    }
                    
        return jsonify({
            'status': 'success',
            'message': f'成功生成分班分組 Excel 檔案。儲存路徑為: {output_path}',
            'stats': response_stats
        })
        
    except Exception as e:
        import traceback
        error_details = traceback.format_exc()
        print(f"Error during class division processing: {e}\n{error_details}")
        return jsonify({'status': 'error', 'message': f'處理分班分組時出錯: {str(e)}', 'details': error_details})

if __name__ == '__main__':
    print("Starting local Flask backend...")
    
    # Run Flask in a daemon thread
    flask_thread = threading.Thread(
        target=lambda: app.run(port=5000, host='127.0.0.1', debug=False),
        daemon=True
    )
    flask_thread.start()
    
    # Auto open browser
    def open_browser():
        webbrowser.open("http://127.0.0.1:5000/")
    threading.Timer(1.5, open_browser).start()
    
    # Main thread runs the Tkinter event loop and handles dialog tasks from Flask thread
    while True:
        try:
            task, res_q = dialog_queue.get(timeout=0.05)
            if task == 'select_folder':
                folder = filedialog.askdirectory(title="選擇考評/分班分組學年主目錄")
                res_q.put(folder)
            elif task == 'select_output':
                file_path = filedialog.asksaveasfilename(
                    title="選擇儲存的 Excel 檔案路徑",
                    defaultextension=".xlsx",
                    filetypes=[("Excel Files", "*.xlsx")]
                )
                res_q.put(file_path)
        except queue.Empty:
            pass
        
        try:
            root.update()
        except tk.TclError:
            # Tkinter root destroyed (exiting)
            break
